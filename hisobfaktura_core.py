#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Hisob-faktura (.xls, aslida HTML) fayllarini SmartPOS savdo saytiga
yuklanadigan "product_list" Excel shabloniga aylantiradigan desktop dastur.

Ustunlar POZITSIYASI bo'yicha emas, balki har bir faylning o'z sarlavha
matniga qarab DINAMIK aniqlanadi — chunki turli ta'minotchilarning hisob-
fakturalarida ustunlar soni farq qilishi mumkin (masalan, markировкasiz
tovarlarda "Маркировка коди" ustuni umuman bo'lmaydi).

Bu fayl GitHub'da saqlanadi va launcher.py orqali har ishga tushganda
avtomatik yangilanadi — bu yerni tahrirlash = barcha foydalanuvchilarning
dasturi keyingi ochilishda yangilanadi degani.
"""
__version__ = "2026-09-08.1"

import os
import re
import sys
import json
import threading
import queue
import datetime
import urllib.request
from concurrent.futures import ThreadPoolExecutor

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from bs4 import BeautifulSoup

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

# ---------------------------------------------------------------------------
# Shablon konstantalari
# ---------------------------------------------------------------------------

HEADERS = [
    "Название товара", "Штрих код", "Ед. измерения", "Цена", "НДС %",
    "Избранное", "Маркировка", "ИКПУ", "Код ед. измерения тасниф",
]

UNITS = [
    "упаковка", "килограмм", "миллилитр", "блок", "коробка", "куб", "ампула",
    "часы", "тюбик", "флакон", "пластиковые банки / бутылки", "проценты",
    "комплект", "куб.м", "тонна", "пачка", "кв. дециметр", "комплект / набор",
    "банка", "мешок", "км", "порция", "сум", "пара", "сотих", "гектар",
    "кв. м.", "грамм", "пг. м.", "сутка", "кВа", "метр", "фут", "блистер",
    "литр", "штук", "человек",
]
VATS = ["Без НДС", "6.00", "0.00", "12.00"]
FAVOURITES = ["Да", "Нет"]

DEFAULT_BARCODE = "1111111111111"  # shtrix kodi yo'q tovarlar uchun (13 ta "1")

# hisob-fakturadagi lotin/o'zbekcha o'lchov birliklarini sayt kutgan
# rus tilidagi nomlarga moslash. Kerak bo'lsa shu ro'yxatga qo'shib boring.
UNIT_MAP = {
    "dona": "штук",
    "kg": "килограмм", "kilogramm": "килограмм",
    "litr": "литр", "l": "литр",
    "metr": "метр", "m": "метр",
    "komplekt": "комплект", "to'plam": "комплект", "toplam": "комплект",
    "quti": "коробка", "karobka": "коробка", "karobka/quti": "коробка",
    "blok": "блок",
    "para": "пара", "juft": "пара",
    "tonna": "тонна",
    "gramm": "грамм", "g": "грамм",
    "flakon": "флакон",
    "tyubik": "тюбик",
    "banka": "банка",
    "paket": "мешок", "meshok": "мешок", "qop": "мешок",
    "blister": "блистер",
    "soat": "часы", "chas": "часы",
    "kub": "куб", "kub.m": "куб.м", "kubm": "куб.м",
    "ampula": "ампула",
    "pachka": "пачка",
    "sum": "сум",
    "km": "км",
    "porsiya": "порция",
    "sutka": "сутка",
    "foiz": "проценты", "protsent": "проценты",
    "inson": "человек", "kishi": "человек", "chelovek": "человек",
    "gektar": "гектар",
    "sotix": "сотих", "sotih": "сотих",
    # kirillcha qisqartma/variantlar — ba'zi hisob-fakturalarda "dona" o'rniga
    # to'g'ridan-to'g'ri shu ko'rinishda keladi
    "шт": "штук", "шт.": "штук", "штук.": "штук", "штука": "штук",
    "кг": "килограмм", "кг.": "килограмм",
    "л": "литр", "л.": "литр",
    "уп": "упаковка", "уп.": "упаковка", "упак": "упаковка", "упак.": "упаковка",
    "компл": "комплект", "компл.": "комплект", "к-т": "комплект",
    "пар": "пара",
    "пач": "пачка", "пач.": "пачка",
}


# ---------------------------------------------------------------------------
# Hisob-fakturani o'qish (dinamik ustun aniqlash)
# ---------------------------------------------------------------------------

def _colspan(cell):
    raw = cell.get("colspan", "1")
    try:
        return max(1, int(float(raw)))
    except (TypeError, ValueError):
        return 1


def _parse_header_row(tr):
    """tr ichidagi har bir katakning MATNI -> boshlang'ich ustun indeksi.

    colspan hisobga olinadi, shu bilan pastdagi data qatorlaridagi haqiqiy
    katak indekslari bilan mos keladi.
    """
    mapping = {}
    col = 0
    for cell in tr.find_all(["td", "th"], recursive=False):
        label = cell.get_text(strip=True)
        if label and label not in mapping:
            mapping[label] = col
        col += _colspan(cell)
    return mapping


def _find_header_map(thead):
    for tr in thead.find_all("tr"):
        text = tr.get_text().lower()
        if "маҳсулот" in text or "наименование" in text:
            return _parse_header_row(tr)
    return None


def _resolve_field_indices(header_map):
    """Sarlavha matnlariga qarab har bir maydonning ustun-indeksini topadi."""
    idx = {}

    def claim(field, col):
        idx.setdefault(field, col)

    for label, col in header_map.items():
        low = label.lower()
        if "маҳсулот" in low or "наименование" in low:
            claim("name", col)
        elif "маркировка" in low:
            claim("marking", col)
        elif "идентификация" in low or "икпу" in low or "каталог" in low:
            claim("ikpu", col)
        elif "штрих" in low:
            claim("barcode", col)
        elif "ўлчов" in low or "измерен" in low:
            claim("unit", col)
        elif "миқдор" in low or "количество" in low:
            claim("qty", col)
        elif "нарҳ" in low or "нарх" in low or low.strip() == "цена":
            claim("price", col)
        elif "ққс" in low or "ндс" in low:
            claim("vat", col)
    return idx


def _find_item_table(soup):
    """Hisob-faktura jadvalini (thead+tbody) qidiradi."""
    for table in soup.find_all("table"):
        thead = table.find("thead")
        tbody = table.find("tbody")
        if thead is None or tbody is None:
            continue
        header_map = _find_header_map(thead)
        if header_map:
            return header_map, tbody
    return None, None


def map_unit(raw, warnings, name):
    base = re.sub(r"\(.*?\)", "", raw).strip().lower()
    if base in UNIT_MAP:
        return UNIT_MAP[base]
    if base in UNITS:
        return base

    # "потребительская коробка=1 шт" kabi tavsifiy matn ichidan ham
    # tanish so'zni (masalan "шт") alohida so'z sifatida qidiramiz —
    # eng uzun kalitdan boshlab, noto'g'ri qisman moslikni oldini olish uchun.
    for key in sorted(UNIT_MAP, key=len, reverse=True):
        if re.search(rf"(?<!\w){re.escape(key)}(?!\w)", base):
            return UNIT_MAP[key]

    mapped = raw.strip()
    if mapped not in UNITS:
        warnings.append(
            f"  - '{name[:60]}': o'lchov birligi '{raw}' saytning ro'yxatida yo'q, "
            f"'{mapped}' deb yozildi — qo'lda tekshiring yoki UNIT_MAP ga qo'shing"
        )
    return mapped


def parse_price(text):
    cleaned = text.replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        value = float(cleaned)
    except ValueError:
        return None
    return format(value, ".2f")  # doim nuqta bilan, vergul bilan emas


def parse_vat(text):
    text = text.strip()
    if "без" in text.lower():
        return "Без НДС"
    m = re.search(r"[\d.,]+", text)
    if not m:
        return text
    num = m.group(0).replace(",", ".")
    try:
        return format(float(num), ".2f")
    except ValueError:
        return text


def extract_ikpu(text):
    m = re.search(r"\d{15,17}", text)
    if m:
        return m.group(0)
    return text.split("\n")[0].strip()


# ---------------------------------------------------------------------------
# ИКПУ -> "Код ед. измерения тасниф" (rasmiy tasnif.soliq.uz orqali)
# ---------------------------------------------------------------------------
# Sayt ИКПУни shu "тасниф" kodi bilan birga tekshiradi — u yo'q bo'lsa,
# hatto to'g'ri ИКПУ ham "ИКПУ неверна" deb rad etiladi.

IKPU_API_URL = "https://tasnif.soliq.uz/api/cls-api/mxik/get/by-mxik"


def fetch_classifier_code(ikpu, timeout=6):
    """Bitta ИКПУ uchun rasmiy 'Код ед. измерения тасниф' kodini oladi.

    Qaytaradi: (kod_yoki_None, xatolik_matni_yoki_None)
    """
    if not ikpu:
        return None, None
    try:
        url = f"{IKPU_API_URL}?mxikCode={ikpu}&lang=uz"
        req = urllib.request.Request(url, headers={"User-Agent": "HisobFaktura2ProductList"})
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            data = json.load(resp)
        packages = data.get("packages") or []
        if not packages:
            return None, "tasnif.soliq.uz'da bu ИКПУ uchun o'lchov birligi (упаковка) topilmadi"
        code = packages[0].get("code")
        return (str(code) if code is not None else None), None
    except Exception as exc:  # noqa: BLE001
        return None, str(exc)


def enrich_classifier_codes(items, cache=None, log=None):
    """Har bir mahsulotning 'classifier_code' maydonini tasnif.soliq.uz'dan
    to'ldiradi. `cache` — {ikpu: kod} lug'ati, bir nechta faylda takrorlangan
    ИКПУлar uchun qayta so'rov yubormaslik uchun (chaqiruvchi saqlab, keyingi
    fayllarga ham shu lug'atni berishi mumkin)."""
    if cache is None:
        cache = {}

    unique_ikpus = sorted({
        it["ikpu"] for it in items
        if it.get("ikpu") and it["ikpu"] not in cache
    })
    if unique_ikpus:
        if log:
            log(f"ИКПУ kodlari tasnif.soliq.uz'dan tekshirilmoqda ({len(unique_ikpus)} ta noyob kod)...")

        def _lookup(code):
            val, err = fetch_classifier_code(code)
            return code, val, err

        with ThreadPoolExecutor(max_workers=8) as ex:
            for code, val, err in ex.map(_lookup, unique_ikpus):
                cache[code] = val
                if err and log:
                    log(f"  DIQQAT: ИКПУ {code} uchun tasnif kodi topilmadi — {err}")

    missing = 0
    for it in items:
        ikpu = it.get("ikpu")
        if ikpu:
            it["classifier_code"] = cache.get(ikpu)
            if not it["classifier_code"]:
                missing += 1
    if missing and log:
        log(f"  {missing} ta mahsulotda 'Код ед. измерения тасниф' topilmadi — saytga yuklashdan oldin qo'lda tekshiring.")
    return cache


# ---------------------------------------------------------------------------
# Foydalanish/xato hisoboti (GitHub Issues orqali kuzatuv)
# ---------------------------------------------------------------------------
# Bu FAQAT dasturning o'z ishlashidagi xatolarni (fayl o'qilmadi, tarmoq
# muammosi va h.k.) kuzatadi — SmartPOS sayti faylni qabul qilgach chiqargan
# xatolar (masalan "ИКПУ неверна") bu yerga tushmaydi, chunki ular dastur
# ishini tugatgandan KEYIN, saytning o'zida yuz beradi.
#
# GitHub'ga YOZISH huquqli token bu yerda SAQLANMAYDI (ochiq fayl bo'lgani
# uchun xavfsiz emas edi — GitHub buni avtomatik bloklab, to'g'ri qildi).
# Buning o'rniga kichik Cloudflare Worker orqali yuboriladi: u tokenni o'zida
# (server tomonida, hech kimga ko'rinmaydigan holda) saqlaydi va shu yerdan
# GitHub'ga o'zi yozadi. Bu yerdagi APP_KEY — GitHub tokeni EMAS, faqat shu
# Worker'ni tasodifiy so'rovlardan himoyalash uchun oddiy ilova kaliti;
# u bilan faqat shu bitta Worker'ga (o'zga hech narsaga) xabar yuborish
# mumkin.

_REPORT_URL = "https://hf2pl-report.tasks-bot.workers.dev"
_REPORT_APP_KEY = "OF28Y-uJR3NxaMKPbQBwNhJhpdeZoNwi"


def report_event(kind, detail=""):
    """Foydalanish/xato hodisasini fon rejimida (bloklamasdan) yuboradi.
    Internet yo'q yoki xato bo'lsa — jim o'tkazib yuboriladi, dastur
    ishlashiga hech qanday ta'sir qilmaydi."""

    def worker():
        try:
            import socket
            host = socket.gethostname()
        except Exception:
            host = "noma'lum"
        payload = {
            "kind": kind,
            "version": __version__,
            "host": host,
            "detail": detail[:1500] if detail else "",
        }
        try:
            data = json.dumps(payload).encode("utf-8")
            req = urllib.request.Request(
                _REPORT_URL, data=data, method="POST",
                headers={
                    "X-App-Key": _REPORT_APP_KEY,
                    "Content-Type": "application/json",
                    "User-Agent": "HisobFaktura2ProductList",
                },
            )
            urllib.request.urlopen(req, timeout=6)
        except Exception:
            pass  # kuzatuv ishlamasa ham, dastur o'z ishini davom ettiradi

    threading.Thread(target=worker, daemon=True).start()


def parse_invoice(path):
    with open(path, encoding="utf-8", errors="ignore") as fh:
        html = fh.read()
    soup = BeautifulSoup(html, "html.parser")
    header_map, tbody = _find_item_table(soup)
    if tbody is None:
        raise RuntimeError(
            "Hisob-faktura jadvali topilmadi. Fayl kutilgan '.xls (HTML)' "
            "hisobvaraq-faktura formatida emas ko'rinadi."
        )

    idx = _resolve_field_indices(header_map)
    missing_required = [f for f in ("name", "price") if f not in idx]
    if missing_required:
        raise RuntimeError(
            "Jadval sarlavhasida quyidagi ustunlar topilmadi: "
            + ", ".join(missing_required)
            + ". Fayl tuzilishi kutilganidan farq qiladi — namuna sifatida yuboring."
        )

    def get(cells, field, default=""):
        i = idx.get(field)
        if i is None or i >= len(cells):
            return default
        return cells[i]

    items = []
    warnings = []
    for tr in tbody.find_all("tr"):
        cells = [c.get_text(strip=True) for c in tr.find_all(["td", "th"])]
        if not cells or not cells[0].strip().isdigit():
            continue  # "Жами" jami qatori yoki bo'sh qator

        name = get(cells, "name")
        if not name:
            continue
        marking_raw = get(cells, "marking")
        ikpu_raw = get(cells, "ikpu")
        barcode = get(cells, "barcode").strip()
        unit_raw = get(cells, "unit")
        price_raw = get(cells, "price")
        vat_raw = get(cells, "vat")

        items.append({
            "name": name,
            "barcode": barcode or DEFAULT_BARCODE,
            "unit": map_unit(unit_raw, warnings, name) if unit_raw else "штук",
            "price": parse_price(price_raw),
            "vat": parse_vat(vat_raw) if vat_raw else "12.00",
            "favourite": "Да",
            "marking": "Да" if marking_raw.strip() == "Маркировкаланган" else "Нет",
            "ikpu": extract_ikpu(ikpu_raw) if ikpu_raw else "",
            "classifier_code": None,
        })

    if not idx.get("unit"):
        warnings.append("  DIQQAT: bu faylda 'Ўлчов бирлиги' ustuni topilmadi, hammasiga 'штук' qo'yildi.")
    if not idx.get("vat"):
        warnings.append("  DIQQAT: bu faylda 'ҚҚС' ustuni topilmadi, hammasiga '12.00' qo'yildi.")
    if not idx.get("ikpu"):
        warnings.append("  DIQQAT: bu faylda ИКПУ ustuni topilmadi — bo'sh qoldirildi, saytga yuklashdan oldin tekshiring.")

    return items, warnings


# ---------------------------------------------------------------------------
# Excel yozish
# ---------------------------------------------------------------------------

def write_product_list(items, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Продукты"
    ws.append(HEADERS)
    for it in items:
        ws.append([
            it["name"], it["barcode"], it["unit"], it["price"], it["vat"],
            it["favourite"], it["marking"], it["ikpu"], it["classifier_code"],
        ])

    ref = wb.create_sheet("Справочник")
    for i, u in enumerate(UNITS, start=1):
        ref.cell(row=i, column=1, value=u)
    for i, v in enumerate(VATS, start=1):
        ref.cell(row=i, column=2, value=v)
    for i, f in enumerate(FAVOURITES, start=1):
        ref.cell(row=i, column=3, value=f)

    wb.defined_names["Units"] = openpyxl.workbook.defined_name.DefinedName(
        "Units", attr_text=f"Справочник!$A$1:$A${len(UNITS)}")
    wb.defined_names["Vats"] = openpyxl.workbook.defined_name.DefinedName(
        "Vats", attr_text=f"Справочник!$B$1:$B${len(VATS)}")
    wb.defined_names["Favourites"] = openpyxl.workbook.defined_name.DefinedName(
        "Favourites", attr_text=f"Справочник!$C$1:$C${len(FAVOURITES)}")

    dv_units = DataValidation(type="list", formula1="=Units", allow_blank=True)
    dv_vats = DataValidation(type="list", formula1="=Vats", allow_blank=True)
    dv_fav = DataValidation(type="list", formula1="=Favourites", allow_blank=True)
    ws.add_data_validation(dv_units)
    ws.add_data_validation(dv_vats)
    ws.add_data_validation(dv_fav)
    dv_units.add("C2:C1001")
    dv_vats.add("E2:E1001")
    dv_fav.add("F2:G1001")

    wb.save(out_path)


def unique_path(out_dir, base_name):
    """base_name.xlsx band bo'lsa _2, _3 ... qo'shib beradi."""
    stem, ext = os.path.splitext(base_name)
    candidate = os.path.join(out_dir, base_name)
    i = 2
    while os.path.exists(candidate):
        candidate = os.path.join(out_dir, f"{stem}_{i}{ext}")
        i += 1
    return candidate


def default_out_name(in_path):
    base = os.path.splitext(os.path.basename(in_path))[0]
    m = re.search(r"(\d+)", base)
    num = m.group(1) if m else base
    return f"product_list_from_invoice_{num}.xlsx"


# ---------------------------------------------------------------------------
# GUI
# ---------------------------------------------------------------------------

# Minimalistik rang palitrasi
BG = "#fafafa"        # oyna foni
CARD = "#ffffff"      # kartochka/log foni
BORDER = "#e5e7eb"    # nozik chiziqlar
TEXT = "#111827"      # asosiy matn
MUTED = "#6b7280"     # ikkinchi darajali matn
ACCENT = "#2563eb"    # asosiy tugma (ko'k)
ACCENT_DARK = "#1d4ed8"
FONT = ("Segoe UI", 10)
FONT_BOLD = ("Segoe UI", 10, "bold")
FONT_TITLE = ("Segoe UI", 13, "bold")
FONT_MONO = ("Consolas", 9)


class App:
    def __init__(self, root):
        self.root = root
        root.title(f"Hisob-faktura → Product list")
        root.geometry("760x830")
        root.minsize(640, 560)
        root.configure(bg=BG)

        self.selected_files = []
        self.output_dir = None
        self.combine_var = tk.BooleanVar(value=True)
        self.log_queue = queue.Queue()
        self.manual_items = []

        self._setup_style()
        self._build_ui()
        self._poll_log_queue()
        self._log(f"Dastur versiyasi: {__version__}")
        if os.environ.get("HF2PL_UPDATE_STATUS"):
            self._log(os.environ["HF2PL_UPDATE_STATUS"])
        report_event("🟢 Ishga tushdi")

        # dastur doim ko'rinadigan, oldingi planda ochilsin — fonda yashirin
        # ishlamasligi uchun.
        root.deiconify()
        root.lift()
        root.attributes("-topmost", True)
        root.after(300, lambda: root.attributes("-topmost", False))
        root.focus_force()

    # ---------- Uslub ----------
    def _setup_style(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass

        style.configure(".", background=BG, foreground=TEXT, font=FONT)
        style.configure("TFrame", background=BG)
        style.configure("Card.TFrame", background=CARD)
        style.configure("TLabel", background=BG, foreground=TEXT, font=FONT)
        style.configure("Muted.TLabel", background=BG, foreground=MUTED, font=FONT)
        style.configure("Title.TLabel", background=BG, foreground=TEXT, font=FONT_TITLE)
        style.configure("Section.TLabel", background=BG, foreground=TEXT, font=FONT_BOLD)
        style.configure("Count.TLabel", background=BG, foreground=MUTED, font=FONT)

        style.configure("TCheckbutton", background=BG, foreground=TEXT, font=FONT)
        style.map("TCheckbutton", background=[("active", BG)])

        style.configure("TSeparator", background=BORDER)

        # Ikkinchi darajali (neytral) tugma
        style.configure(
            "Secondary.TButton", background=CARD, foreground=TEXT,
            font=FONT, borderwidth=1, relief="solid", padding=(12, 7),
        )
        style.map(
            "Secondary.TButton",
            background=[("active", "#f3f4f6"), ("disabled", "#f3f4f6")],
            foreground=[("disabled", MUTED)],
            bordercolor=[("!disabled", BORDER)],
        )

        # Asosiy (accent) tugma
        style.configure(
            "Primary.TButton", background=ACCENT, foreground="#ffffff",
            font=FONT_BOLD, borderwidth=0, padding=(16, 9),
        )
        style.map(
            "Primary.TButton",
            background=[("active", ACCENT_DARK), ("disabled", "#9ca3af")],
            foreground=[("disabled", "#f3f4f6")],
        )

    # ---------- UI qurish ----------
    def _section_title(self, parent, text):
        ttk.Label(parent, text=text, style="Section.TLabel").pack(anchor="w", pady=(0, 8))

    def _separator(self, parent):
        ttk.Separator(parent, orient="horizontal").pack(fill="x", pady=13)

    def _build_ui(self):
        # Tashqi konteyner: pastki "footer" (Boshlash tugmasi) DOIM ko'rinadi —
        # oyna kichraytirilsa, o'sib-kichrayadigan qism (jurnal) siqiladi,
        # tugmalar hech qachon kesilib qolmaydi.
        root_container = ttk.Frame(self.root)
        root_container.pack(fill="both", expand=True)

        footer = ttk.Frame(root_container, padding=(24, 0, 24, 20))
        footer.pack(side="bottom", fill="x")

        outer = ttk.Frame(root_container, padding=(24, 24, 24, 0))
        outer.pack(side="top", fill="both", expand=True)

        title_label = ttk.Label(outer, text="Hisob-faktura → Product list", style="Title.TLabel")
        title_label.pack(anchor="w")
        subtitle_label = ttk.Label(
            outer, text="Hisob-faktura fayllarini SmartPOS uchun tayyor Excelga aylantiring",
            style="Muted.TLabel",
        )
        subtitle_label.pack(anchor="w", pady=(2, 0))

        # Oyna kengligiga qarab sarlavha ostidagi matn o'ralishi
        def _on_resize(event):
            if event.widget is self.root:
                subtitle_label.configure(wraplength=max(300, event.width - 48))
        self.root.bind("<Configure>", _on_resize)

        self._separator(outer)

        # 1) Fayllar
        self._section_title(outer, "1. Fayllarni tanlang")
        btn_row = ttk.Frame(outer)
        btn_row.pack(fill="x")
        ttk.Button(btn_row, text="Fayl tanlash", style="Primary.TButton",
                   command=self.choose_files).pack(side="left")
        ttk.Button(btn_row, text="Tozalash", style="Secondary.TButton",
                   command=self.clear_files).pack(side="left", padx=(8, 0))
        self.files_count_label = ttk.Label(btn_row, text="Fayl tanlanmagan", style="Count.TLabel")
        self.files_count_label.pack(side="left", padx=(14, 0))

        list_wrap = tk.Frame(outer, bg=BORDER)
        list_wrap.pack(fill="x", pady=(12, 0))
        self.files_listbox = tk.Listbox(
            list_wrap, height=4, selectmode="extended", relief="flat",
            bg=CARD, fg=TEXT, font=FONT, highlightthickness=0,
            selectbackground=ACCENT, selectforeground="#ffffff",
        )
        self.files_listbox.pack(fill="x", padx=1, pady=1)

        self._separator(outer)

        # 2) Saqlash
        self._section_title(outer, "2. Natijani qayerga saqlash")
        row2 = ttk.Frame(outer)
        row2.pack(fill="x")
        ttk.Button(row2, text="Papka tanlash", style="Secondary.TButton",
                   command=self.choose_output_dir).pack(side="left")
        self.output_label = ttk.Label(row2, text="Fayl turgan joyga saqlanadi", style="Muted.TLabel")
        self.output_label.pack(side="left", padx=(14, 0))

        ttk.Checkbutton(
            outer, text="Barchasini BITTA Excel fayliga birlashtirish",
            variable=self.combine_var,
        ).pack(anchor="w", pady=(14, 0))

        self._separator(outer)

        # 3) Qo'lda qo'shish
        self._section_title(outer, "3. Yoki mahsulotni qo'lda kiriting")
        ttk.Label(
            outer,
            text="Fayl orqali to'g'ri o'qilmaydigan yoki alohida tovar uchun.",
            style="Muted.TLabel",
        ).pack(anchor="w", pady=(0, 8))
        row_manual = ttk.Frame(outer)
        row_manual.pack(fill="x")
        ttk.Button(row_manual, text="Mahsulot qo'shish", style="Secondary.TButton",
                   command=self.open_manual_entry_dialog).pack(side="left")
        self.manual_count_label = ttk.Label(row_manual, text="Qo'lda qo'shilgan yo'q", style="Muted.TLabel")
        self.manual_count_label.pack(side="left", padx=(14, 0))

        # 4) Boshlash — doim ko'rinadigan footer'da
        self._separator(footer)
        step3 = ttk.Frame(footer)
        step3.pack(fill="x")
        self.start_btn = ttk.Button(step3, text="Boshlash", style="Primary.TButton",
                                     command=self.start_conversion)
        self.start_btn.pack(side="left")
        self.open_folder_btn = ttk.Button(
            step3, text="Papkani ochish", style="Secondary.TButton",
            command=self.open_output_folder, state="disabled",
        )
        self.open_folder_btn.pack(side="left", padx=(8, 0))

        self._separator(outer)

        # Jurnal / natija
        ttk.Label(outer, text="Jarayon", style="Muted.TLabel").pack(anchor="w", pady=(20, 6))
        log_wrap = tk.Frame(outer, bg=BORDER)
        log_wrap.pack(fill="both", expand=True)
        log_inner = tk.Frame(log_wrap, bg=CARD)
        log_inner.pack(fill="both", expand=True, padx=1, pady=1)
        self.log_text = tk.Text(
            log_inner, wrap="word", state="disabled", relief="flat",
            bg=CARD, fg=TEXT, font=FONT_MONO, highlightthickness=0,
            padx=10, pady=8, height=8,
        )
        scroll = ttk.Scrollbar(log_inner, command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=scroll.set)
        self.log_text.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")

    # ---------- Fayl tanlash ----------
    def choose_files(self):
        paths = filedialog.askopenfilenames(
            title="Hisob-faktura fayllarini tanlang (bir nechtasini birga tanlash mumkin)",
            filetypes=[("Hisob-faktura fayllari", "*.xls *.xlsx *.html *.htm"),
                       ("Barcha fayllar", "*.*")],
        )
        for p in paths:
            if p not in self.selected_files:
                self.selected_files.append(p)
                self.files_listbox.insert("end", os.path.basename(p))
        self._update_files_count_label()

    def _update_files_count_label(self):
        n = len(self.selected_files)
        text = "Fayl tanlanmagan" if n == 0 else f"{n} ta fayl tanlandi"
        self.files_count_label.config(text=text)

    def clear_files(self):
        self.selected_files = []
        self.files_listbox.delete(0, "end")
        self._update_files_count_label()

    # ---------- Qo'lda kiritish ----------
    def _update_manual_count_label(self):
        n = len(self.manual_items)
        text = "Qo'lda qo'shilgan yo'q" if n == 0 else f"{n} ta qo'lda qo'shildi"
        self.manual_count_label.config(text=text)

    def open_manual_entry_dialog(self):
        win = tk.Toplevel(self.root)
        win.title("Mahsulot qo'shish")
        win.configure(bg=BG)
        win.geometry("460x600")
        win.minsize(420, 560)
        win.transient(self.root)
        win.grab_set()

        form = ttk.Frame(win, padding=20)
        form.pack(fill="both", expand=True)

        def field(label_text):
            ttk.Label(form, text=label_text, style="TLabel").pack(anchor="w", pady=(10, 3))

        field("Nomi *")
        name_entry = ttk.Entry(form, font=FONT)
        name_entry.pack(fill="x")

        field("Shtrix kod (bo'sh — avtomatik 1111111111111)")
        barcode_entry = ttk.Entry(form, font=FONT)
        barcode_entry.pack(fill="x")

        field("Ед. измерения")
        unit_combo = ttk.Combobox(form, values=UNITS, state="readonly", font=FONT)
        unit_combo.set("штук")
        unit_combo.pack(fill="x")

        field("Narx (so'm) *")
        price_entry = ttk.Entry(form, font=FONT)
        price_entry.pack(fill="x")

        field("НДС %")
        vat_combo = ttk.Combobox(form, values=VATS, state="readonly", font=FONT)
        vat_combo.set("12.00")
        vat_combo.pack(fill="x")

        field("ИКПУ (17 xonali kod)")
        ikpu_row = ttk.Frame(form)
        ikpu_row.pack(fill="x")
        ikpu_entry = ttk.Entry(ikpu_row, font=FONT)
        ikpu_entry.pack(side="left", fill="x", expand=True)
        classifier_status = ttk.Label(form, text="", style="Muted.TLabel")

        classifier_code_holder = {"value": None}

        def lookup_classifier():
            ikpu = ikpu_entry.get().strip()
            if not ikpu:
                return
            classifier_status.config(text="Qidirilmoqda...")

            def worker():
                code, err = fetch_classifier_code(ikpu)
                def apply():
                    classifier_code_holder["value"] = code
                    if code:
                        classifier_status.config(text=f"✓ Tasnif kodi topildi: {code}")
                    else:
                        classifier_status.config(text=f"⚠ Topilmadi: {err or 'ИКПУ noto\'g\'ri bo\'lishi mumkin'}")
                self.root.after(0, apply)

            threading.Thread(target=worker, daemon=True).start()

        ttk.Button(ikpu_row, text="Tasnif kodini top", style="Secondary.TButton",
                   command=lookup_classifier).pack(side="left", padx=(8, 0))
        classifier_status.pack(anchor="w", pady=(4, 0))

        vars_row = ttk.Frame(form)
        vars_row.pack(fill="x", pady=(14, 0))
        favourite_var = tk.BooleanVar(value=True)
        marking_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(vars_row, text="Избранное", variable=favourite_var).pack(side="left")
        ttk.Checkbutton(vars_row, text="Маркировка", variable=marking_var).pack(side="left", padx=(20, 0))

        error_label = ttk.Label(form, text="", style="Muted.TLabel", foreground="#dc2626")
        error_label.pack(anchor="w", pady=(10, 0))

        btn_row = ttk.Frame(form)
        btn_row.pack(fill="x", pady=(16, 0))

        def submit(close_after):
            name = name_entry.get().strip()
            price_raw = price_entry.get().strip().replace(",", ".")
            if not name or not price_raw:
                error_label.config(text="Nomi va narxni to'ldiring.")
                return
            try:
                price = format(float(price_raw), ".2f")
            except ValueError:
                error_label.config(text="Narx noto'g'ri formatda.")
                return

            barcode = barcode_entry.get().strip() or DEFAULT_BARCODE
            ikpu = ikpu_entry.get().strip()
            item = {
                "name": name,
                "barcode": barcode,
                "unit": unit_combo.get() or "штук",
                "price": price,
                "vat": vat_combo.get() or "12.00",
                "favourite": "Да" if favourite_var.get() else "Нет",
                "marking": "Да" if marking_var.get() else "Нет",
                "ikpu": ikpu,
                "classifier_code": classifier_code_holder["value"],
            }
            self.manual_items.append(item)
            self._update_manual_count_label()
            self._log(f"Qo'lda qo'shildi: {name}")

            if close_after:
                win.destroy()
            else:
                name_entry.delete(0, "end")
                barcode_entry.delete(0, "end")
                price_entry.delete(0, "end")
                ikpu_entry.delete(0, "end")
                classifier_status.config(text="")
                classifier_code_holder["value"] = None
                error_label.config(text="")
                name_entry.focus_set()

        ttk.Button(btn_row, text="Qo'shish va yana", style="Secondary.TButton",
                   command=lambda: submit(False)).pack(side="left")
        ttk.Button(btn_row, text="Qo'shish va yopish", style="Primary.TButton",
                   command=lambda: submit(True)).pack(side="left", padx=(8, 0))

        name_entry.focus_set()

    def choose_output_dir(self):
        d = filedialog.askdirectory(title="Natijalarni qayerga saqlash kerak?")
        if d:
            self.output_dir = d
            self.output_label.config(text=d)
            self._log(f"Saqlash papkasi tanlandi: {d}")
        else:
            self.output_dir = None
            self.output_label.config(text="Fayl turgan joyga saqlanadi")

    def open_output_folder(self):
        d = self.output_dir or (os.path.dirname(self.selected_files[0]) if self.selected_files else None)
        if d and os.path.isdir(d):
            os.startfile(d)

    def _default_out_dir(self):
        if self.output_dir:
            return self.output_dir
        if self.selected_files:
            return os.path.dirname(self.selected_files[0]) or "."
        desktop = os.path.join(os.path.expanduser("~"), "Desktop")
        return desktop if os.path.isdir(desktop) else os.path.expanduser("~")

    # ---------- Log ----------
    def _log(self, text):
        self.log_queue.put(text)

    def _poll_log_queue(self):
        try:
            while True:
                text = self.log_queue.get_nowait()
                self.log_text.configure(state="normal")
                self.log_text.insert("end", text + "\n")
                self.log_text.see("end")
                self.log_text.configure(state="disabled")
        except queue.Empty:
            pass
        self.root.after(150, self._poll_log_queue)

    # ---------- Konvertatsiya ----------
    def start_conversion(self):
        if not self.selected_files and not self.manual_items:
            messagebox.showwarning(
                "Diqqat",
                "Avval hisob-faktura fayl(lar)ini tanlang yoki 'Mahsulot qo'shish' orqali qo'lda kiriting.",
            )
            return
        self.start_btn.config(state="disabled", text="Ishlanmoqda...")
        self.open_folder_btn.config(state="disabled")
        thread = threading.Thread(target=self._run_conversion, daemon=True)
        thread.start()

    def _run_conversion(self):
        combine = self.combine_var.get()
        combined_items = []
        saved_paths = []
        error_count = 0
        classifier_cache = {}

        if self.manual_items:
            self._log(f"\n--- Qo'lda kiritilgan mahsulotlar ({len(self.manual_items)} ta) ---")
            try:
                enrich_classifier_codes(self.manual_items, classifier_cache, self._log)
                if combine:
                    combined_items.extend(self.manual_items)
                else:
                    out_dir = self._default_out_dir()
                    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
                    out_path = unique_path(out_dir, f"product_list_qolda_{stamp}.xlsx")
                    write_product_list(self.manual_items, out_path)
                    saved_paths.append(out_path)
                    self._log(f"Saqlandi: {out_path}")
            except Exception as exc:  # noqa: BLE001
                error_count += 1
                self._log(f"XATOLIK (qo'lda kiritilganlarni saqlashda): {exc}")
                report_event("🔴 Xato (qo'lda kiritishda)", str(exc))

        for path in self.selected_files:
            base = os.path.basename(path)
            self._log(f"\n--- {base} ---")
            try:
                self._log("Fayl o'qilmoqda...")
                items, warnings = parse_invoice(path)
                self._log(f"{len(items)} ta mahsulot topildi.")
                for w in warnings:
                    self._log(w)

                enrich_classifier_codes(items, classifier_cache, self._log)

                if combine:
                    combined_items.extend(items)
                else:
                    out_dir = self.output_dir or os.path.dirname(path) or "."
                    out_path = unique_path(out_dir, default_out_name(path))
                    write_product_list(items, out_path)
                    saved_paths.append(out_path)
                    self._log(f"Saqlandi: {out_path}")
            except Exception as exc:  # noqa: BLE001
                error_count += 1
                self._log(f"XATOLIK: {exc}")
                report_event("🔴 Xato (fayl o'qishda)", f"fayl: {base}\nxato: {exc}")

        if combine and combined_items:
            try:
                out_dir = self._default_out_dir()
                stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
                out_path = unique_path(out_dir, f"product_list_combined_{stamp}.xlsx")
                write_product_list(combined_items, out_path)
                saved_paths.append(out_path)
                self._log(f"\n=== Barchasi birlashtirildi: {len(combined_items)} ta mahsulot ===")
                self._log(f"Saqlandi: {out_path}")
            except Exception as exc:  # noqa: BLE001
                error_count += 1
                self._log(f"XATOLIK (saqlashda): {exc}")
                report_event("🔴 Xato (birlashtirib saqlashda)", str(exc))

        self._log(f"\nTayyor. {len(self.selected_files)} ta fayl ishlandi"
                   f"{f', {error_count} tasida xatolik' if error_count else ''}.")

        self.root.after(0, self._on_conversion_done, saved_paths)

    def _on_conversion_done(self, saved_paths):
        self.start_btn.config(state="normal", text="Boshlash")
        if saved_paths:
            self.open_folder_btn.config(state="normal")
        self.root.lift()
        messagebox.showinfo(
            "Tayyor",
            f"Konvertatsiya tugadi.\n\nSaqlangan fayllar soni: {len(saved_paths)}"
            + ("\n\n" + "\n".join(saved_paths[:10]) if saved_paths else ""),
        )


def _enable_dpi_awareness():
    """Tk() yaratilishidan OLDIN chaqiriladi — Windows monitor masshtabi
    125%/150% bo'lganda elementlar siljib ustma-ust tushib qolmasligi uchun."""
    try:
        from ctypes import windll
        windll.shcore.SetProcessDpiAwareness(1)  # PROCESS_SYSTEM_DPI_AWARE
    except Exception:
        try:
            from ctypes import windll
            windll.user32.SetProcessDPIAware()  # eski Windows uchun
        except Exception:
            pass


def main():
    _enable_dpi_awareness()
    root = tk.Tk()
    try:
        dpi = root.winfo_fpixels("1i")
        if dpi > 0:
            root.tk.call("tk", "scaling", dpi / 72.0)
    except Exception:
        pass
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
